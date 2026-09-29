#!/usr/bin/env python3
"""
anr_export.py: export an Obsidian-flavoured Markdown grant draft to an ANR
AAPG pre-proposal, laid out like the official Word template.

    python anr_export.py Proposal.md            # -> Proposal.pdf  (via LaTeX, default)
    python anr_export.py Proposal.md --docx     # -> Proposal.docx (filled Word template)

The LaTeX sources and converted figures are kept in a hidden folder next to
the output (".<name>_build/"), so you can inspect or tweak the .tex.

Requirements
------------
    pandoc                      brew install pandoc
    a TeX distribution          MacTeX / BasicTeX (xelatex)
    SVG figures                 brew install librsvg   (or: pip install cairosvg)
    carlito font                brew install --cask font-carlito
    --docx only                 pip install python-docx lxml

Markdown conventions
--------------------
* %% ... %%                 Obsidian comments: removed (inline or multi-line).
* Preamble "Key: value" lines, before the first heading, fill the header:
      Title, Acronym, Duration, Instrument, Coordinated by, Scientific Theme,
      Bibliography File (.bib), and optionally Margin (e.g. 2cm) and
      Figure Width (figure boxes relative to the text width, default 1.2:
      the box sticks out by 10% of the text width into each margin).
* ## / ### / ####           Top heading level used -> section (I., II., ...),
                            next -> subsection (I.1, I.2, ...), next -> unnumbered.
                            Manual numbers such as "1." are stripped.
* [Author et al., 2015](Key.pdf)
                            Numbered citation "[1]", linked to the reference list;
                            Key is looked up in the .bib file. Numbers follow the
                            order of first citation. "([A, 2015](A.pdf); [B, 2016](B.pdf))"
                            -> "[1, 2]"; "[Isen and colleagues (1987)](Isen1987.pdf)"
                            -> "Isen and colleagues [3]".
* Reference list: one running paragraph,
      [1] Author et al., *Journal* (Year) Title.
  where "Author et al., Journal (Year)" links to the DOI.
* %%beginFigure%% ... %%endFigure%%
      %% location: top %%   optional: force the box to the top (or bottom) of a page
      ![](path/fig.svg)     image, at its own size; optionally {side=right}
      **| Short title.** …  legend; numbered automatically ("Fig. 2 | ...")
                            PDF: framed floating box; the legend starts beside
                            the image and continues below it (goes fully below
                            if less than 4 cm is left beside the image).
* [Fig.](path/fig.svg), [Fig.b](path/fig.svg)
                            Cross-reference -> "Fig. 2", "Fig. 2b".
* An empty "References" heading receives the bibliography (added at the end
  if there is no such heading).
"""
from __future__ import annotations

import argparse
import os
import re
import shutil
import json
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = HERE / "anr-aapg-2027-template.docx"

PANDOC_FORMAT = ("markdown"
                 "-blank_before_header"
                 "+lists_without_preceding_blankline"
                 "+implicit_figures"
                 "-auto_identifiers")

META_KEYS = {
    "title": "title", "acronym": "acronym", "duration": "duration",
    "instrument": "instrument", "coordinated by": "coordinator",
    "coordinator": "coordinator", "scientific theme": "theme", "theme": "theme",
    "bibliography file": "bibliography", "bibliography": "bibliography",
    "margin": "margin", "margins": "margin",
    "figure width": "figwidth",
}

IMG_EXT = r"(?:svg|png|jpe?g|gif|tiff?|bmp|pdf|eps)"
FIG_BLOCK = re.compile(r"%%\s*beginFigure\s*%%(.*?)%%\s*endFigure\s*%%", re.S)
FIG_MARKER = re.compile(r"%%\s*(?:beginFigure|endFigure|location\s*:\s*\w+)\s*%%", re.I)
PLACEMENT = {"top": "t", "bottom": "b", "here": "h", "page": "p"}

warnings: list[str] = []


def warn(msg: str) -> None:
    if msg not in warnings:
        warnings.append(msg)


# ==========================================================================
# 1. Markdown pre-processing (shared by the PDF and DOCX targets)
# ==========================================================================

def strip_comments(text: str, keep_figure_markers: bool = False) -> str:
    """Remove %% ... %% comments (optionally keeping the figure-block markers)."""
    def sub(m):
        return m.group(0) if keep_figure_markers and FIG_MARKER.fullmatch(m.group(0)) else ""
    return re.sub(r"%%.*?%%", sub, text, flags=re.S)


def extract_preamble(text: str) -> tuple[dict, str]:
    """Pull 'Key: value' lines (known keys only) from before the first heading."""
    meta: dict[str, str] = {}
    out, in_preamble = [], True
    for line in text.splitlines():
        if in_preamble and re.match(r"#{1,6}\s", line):
            in_preamble = False
        if in_preamble:
            m = re.match(r"^\s*([A-Za-z][A-Za-z ]{0,30}?)\s*:\s*(.+?)\s*$", line)
            if m and m.group(1).strip().lower() in META_KEYS:
                meta[META_KEYS[m.group(1).strip().lower()]] = m.group(2)
                continue
        out.append(line)
    return meta, "\n".join(out)


def file_key(path: str) -> str:
    return os.path.basename(path.strip().strip("<>")).lower()


def fig_label(path: str) -> str:
    return "fig:" + re.sub(r"[^A-Za-z0-9]+", "-", Path(path.strip().strip("<>")).stem).strip("-")


def scan_figures(text: str) -> dict[str, tuple[int, str]]:
    """Figure blocks in order of appearance: file name -> (number, label)."""
    figs: dict[str, tuple[int, str]] = {}
    for n, block in enumerate(FIG_BLOCK.finditer(text), 1):
        m = re.search(r"!\[[^\]]*\]\(([^)]+)\)|!\[\[([^\]|]+)", block.group(1))
        if m:
            p = m.group(1) or m.group(2)
            figs[file_key(p)] = (n, fig_label(p))
        else:
            warn(f"Figure block {n} contains no image.")
    return figs


def replace_figure_refs(text: str, figs, target: str) -> str:
    pat = re.compile(r"(?<!!)\[\s*(Fig(?:ure)?s?\.?)\s*([^\]]*)\]\(([^)]+\." + IMG_EXT + r")\)", re.I)

    def sub(m):
        label, suffix, dest = m.group(1), m.group(2).strip(), m.group(3)
        if not label.endswith(".") and not label.lower().startswith("figure"):
            label += "."
        fig = figs.get(file_key(dest))
        if fig is None:
            warn(f"Cross-reference to a figure that is not in the document: {dest}")
            return f"{label} ??{suffix}"
        if target == "latex":
            return f"`{label}~\\ref{{{fig[1]}}}{suffix}`{{=latex}}"
        return f"{label} {fig[0]}{suffix}"

    return pat.sub(sub, text)


CITE_TOKEN = "\x00C{}\x00"


def replace_citations(text: str, known_keys: set[str], target: str) -> tuple[str, list[str]]:
    """Numbered citations, in order of first appearance.

    [Author et al., 2015](Key.pdf)          -> [n]
    [Isen and colleagues (1987)](Key.pdf)   -> Isen and colleagues [n]
    ([A, 2015](A.pdf); [B, 2016](B.pdf))   -> [n, m]
    Keys missing from the .bib keep their text as written.
    """
    keys: list[str] = []
    missing: list[str] = []
    pat = re.compile(r"(?<!!)\[([^\]]+)\]\(<?([^()\s/\\<>]+)\.pdf>?\)")

    def cite(m):
        txt, key = m.group(1).strip(), m.group(2)
        if key not in known_keys:
            missing.append(key)
            return txt
        if key not in keys:
            keys.append(key)
        token = CITE_TOKEN.format(keys.index(key) + 1)
        narrative = re.match(r"^(.*?)\s*\(\s*\d{4}[a-z]?\s*\)$", txt)
        if narrative:                                   # "Isen and colleagues (1987)"
            return f"{narrative.group(1)} {token}"
        if re.search(r"\d{4}[a-z]?$", txt):             # "Author et al., 2015"
            return token
        return f"{txt} {token}"                          # free text

    text = pat.sub(cite, text)
    if missing and known_keys:
        warn("Not found in the .bib (check the PDF file names / citation keys): "
             + ", ".join(dict.fromkeys(missing)))

    tok = r"\x00C(\d+)\x00"
    group = rf"{tok}(?:\s*[;,]\s*{tok})*"
    # "(tok; tok)" -> one bracket group, parentheses dropped
    text = re.sub(rf"\(\s*({group})\s*\)", lambda m: render_group(m.group(1), keys, target), text)
    # remaining runs of tokens ("tok; tok" or a single tok)
    text = re.sub(group, lambda m: render_group(m.group(0), keys, target), text)
    return text, keys


def colour(md: str, target: str) -> str:
    """Colour a piece of markdown in the PDF (the Word output keeps its link style)."""
    return f"`{{\\color{{anrblue}}`{{=latex}}{md}`}}`{{=latex}}" if target == "latex" else md


def render_group(run: str, keys: list[str], target: str) -> str:
    nums = sorted(set(int(n) for n in re.findall(r"\x00C(\d+)\x00", run)))
    link = lambda n: f"[{colour(str(n), target)}](#ref-{keys[n - 1]})"
    parts, i = [], 0
    while i < len(nums):                     # compress 3+ consecutive numbers: 2–4
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        if j - i >= 2:
            parts.append(f"{link(nums[i])}\u2013{link(nums[j])}")
        else:
            parts += [link(n) for n in nums[i:j + 1]]
        i = j + 1
    return "\\[" + ", ".join(parts) + "\\]"


def md_escape(s: str) -> str:
    return re.sub(r"([\\`*_{}\[\]<>#$@^~|])", r"\\\1", s)


def load_bib(bib: Path, build: Path) -> dict[str, dict]:
    """Parse the .bib with pandoc (handles LaTeX accents, braces...), cached in the build dir."""
    cache = build / "bib.json"
    if cache.exists() and cache.stat().st_mtime >= bib.stat().st_mtime:
        entries = json.loads(cache.read_text(encoding="utf-8"))
    else:
        r = subprocess.run([find_pandoc(), str(bib), "-f", "bibtex", "-t", "csljson"],
                           capture_output=True, text=True)
        if r.returncode != 0:
            warn("Could not parse the .bib file with pandoc:\n" + r.stderr[-1000:])
            return {}
        cache.write_text(r.stdout, encoding="utf-8")
        entries = json.loads(r.stdout)
    return {e["id"]: e for e in entries if "id" in e}


def format_reference(n: int, key: str, e: dict, target: str) -> str:
    """[n] Author et al., *Journal* (Year) Title.   ('Author ... (Year)' links to the DOI)"""
    def name(a):
        if "literal" in a:
            return a["literal"]
        return " ".join(x for x in (a.get("non-dropping-particle"), a.get("family")) if x)

    authors = [name(a) for a in e.get("author", e.get("editor", []))]
    if not authors:
        who = ""
    elif len(authors) == 1:
        who = authors[0]
    elif len(authors) == 2:
        who = f"{authors[0]} & {authors[1]}"
    else:
        who = f"{authors[0]} et al."
    journal = e.get("container-title") or e.get("publisher") or ""
    try:
        year = str(e["issued"]["date-parts"][0][0])
    except (KeyError, IndexError, TypeError):
        year = "n.d."
    title = re.sub(r"<[^>]+>", "", e.get("title", "")).strip()
    if title and title[-1] not in ".?!":
        title += "."

    head = md_escape(who) + (", " if who and journal else "")
    head += f"*{md_escape(journal)}*" if journal else ""
    head += f" ({year})"
    doi = e.get("DOI", "").strip()
    url = ("https://doi.org/" + re.sub(r"^https?://(dx\.)?doi\.org/", "", doi)) if doi else e.get("URL", "")
    head = colour(head, target)
    if url:
        head = f"[{head}](<{url}>)"
    return f"[**\\[{n}\\]**]{{#ref-{key}}} {head} {md_escape(title)}"


def bibliography_block(keys: list[str], bib: dict, target: str) -> str:
    para = " ".join(format_reference(i + 1, k, bib[k], target) for i, k in enumerate(keys))
    if target == "latex":
        return f"\n\n```{{=latex}}\n\\begin{{anrrefs}}\n```\n\n{para}\n\n```{{=latex}}\n\\end{{anrrefs}}\n```\n"
    return f"\n\n::: {{custom-style=\"Bibliography\"}}\n{para}\n:::\n"


def build_figures(text: str, md_dir: Path, build: Path, target: str) -> str:
    counter = [0]

    def sub(m):
        counter[0] += 1
        n = counter[0]
        img, attrs, caption, place = None, "", [], ""
        for line in (l.strip() for l in m.group(1).strip().splitlines()):
            loc = re.fullmatch(r"%%\s*location\s*:\s*(\w+)\s*%%", line, re.I)
            if loc:
                place = PLACEMENT.get(loc.group(1).lower(), "")
                if not place:
                    warn(f"Figure {n}: unknown location '{loc.group(1)}' (use top, bottom, here or page).")
                continue
            if not line:
                continue
            im = re.match(r"!\[[^\]]*\]\(([^)]+)\)\s*(\{[^}]*\})?\s*$", line)
            wk = re.match(r"!\[\[([^\]|]+)(?:\|[^\]]*)?\]\]\s*(\{[^}]*\})?\s*$", line)
            if img is None and (im or wk):
                g = im or wk
                img, attrs = g.group(1).strip().strip("<>"), (g.group(2) or "")
            else:
                caption.append(line)
        if img is None:
            return ""
        cap = " ".join(caption)
        head = re.match(r"\*\*\s*\|\s*", cap)
        path = prepare_image(img, md_dir, build, target)
        inner = attrs.strip("{} ")

        if target == "latex":
            return latex_figure_box(path, fig_label(img), inner,
                                    ("**" + cap[head.end():]) if head else cap, place)

        # Word: write the number ourselves, pandoc figure with caption
        cap = (f"**Fig. {n} | " + cap[head.end():]) if head else (f"**Fig. {n}.** " + cap).strip()
        cap = cap.replace("[", r"\[").replace("]", r"\]")
        inner = re.sub(r"side\s*=\s*\w+", "", inner).strip()   # natural size unless width given
        return f"\n\n![{cap}](<{path}>){{{inner}}}\n\n"

    return FIG_BLOCK.sub(sub, text)


def latex_figure_box(path: str, label: str, attrs: str, caption: str, place: str = "") -> str:
    """Framed floating box: the image keeps its own size (so the font sizes set
    in the figure are preserved) and the legend starts beside it, then wraps
    below it at the full width of the box.

    Optional attributes after the image, e.g. {side=right}:
      side   left (default) or right.
      width  only to override the natural size (e.g. width=8cm or width=50%).
    place: LaTeX float placement forced by a "%% location: top %%" line ("t", "b"...).
    If the image is wider than the box it is scaled down to fit; if less than
    4 cm is left beside it, the legend goes below the image.
    """
    side = "right" if re.search(r"side\s*=\s*r", attrs) else "left"
    w = re.search(r"width\s*=\s*([\d.]+)\s*(%|cm|mm|in|pt)?", attrs)
    if not w:
        width = ""
    elif (w.group(2) or "%") == "%" and float(w.group(1)) > 1.5:
        width = f"{float(w.group(1)) / 100:.3f}\\anrinnerw"
    elif w.group(2) in (None, "%"):
        width = f"{float(w.group(1)):.3f}\\anrinnerw"
    else:
        width = w.group(1) + w.group(2)
    tex_path = path.replace("\\", "/")
    begin = (f"\\anrsetfig{{{tex_path}}}{{{width}}}\n"
             f"\\begin{{anrfigbox}}[height from={{\\anrfigminheight}} to \\textheight"
             + (f", float={place}" if place else "") + "]")
    legend = f"`\\anrwrap{side}\\anrfignum{{{label}}}`{{=latex}}{caption}"
    return (f"\n\n```{{=latex}}\n{begin}\n```\n\n{legend}\n\n"
            f"```{{=latex}}\n\\anrfigend\n\\end{{anrfigbox}}\n```\n\n")


def prepare_image(img: str, md_dir: Path, build: Path, target: str) -> str:
    src = Path(os.path.expanduser(img))
    if not src.is_absolute():
        src = (md_dir / src).resolve()
    if not src.exists():
        warn(f"Image not found: {img}")
        return str(src)
    ext = src.suffix.lower()
    if ext == ".svg":
        fmt = "pdf" if target == "latex" else "png"
        dst = build / "figures" / f"{src.stem}.{fmt}"
        dst.parent.mkdir(parents=True, exist_ok=True)
        if dst.exists() and dst.stat().st_mtime >= src.stat().st_mtime:
            return str(dst)
        if convert_svg(src, dst, fmt):
            return str(dst)
        if target == "latex":
            sys.exit(f"Cannot convert {src.name} to PDF. Install librsvg "
                     "(`brew install librsvg`) or cairosvg (`pip install cairosvg`).")
        warn(f"No SVG converter found; {src.name} embedded as SVG (fine in recent Word).")
    elif ext == ".pdf" and target == "docx":
        warn(f"{src.name}: PDF figures cannot go into Word; export it as PNG or SVG.")
    return str(src)


def convert_svg(src: Path, dst: Path, fmt: str) -> bool:
    if shutil.which("rsvg-convert"):
        dpi = ["-d", "300", "-p", "300"] if fmt == "png" else []   # PDF: keep the SVG's own size
        r = subprocess.run(["rsvg-convert", "-f", fmt, *dpi, "-o", str(dst), str(src)],
                           capture_output=True)
        if r.returncode == 0:
            return True
    try:
        import cairosvg  # type: ignore
        if fmt == "pdf":
            cairosvg.svg2pdf(url=str(src), write_to=str(dst))
        else:
            cairosvg.svg2png(url=str(src), write_to=str(dst), dpi=300, scale=300 / 96)
        return True
    except Exception:
        pass
    if shutil.which("inkscape"):
        r = subprocess.run(["inkscape", str(src), f"--export-type={fmt}", "--export-dpi=300",
                            f"--export-filename={dst}"], capture_output=True)
        return r.returncode == 0
    return False


def normalise_headings(text: str) -> str:
    levels = [len(m.group(1)) for m in re.finditer(r"^(#{1,6})\s", text, re.M)]
    if not levels:
        return text
    shift = min(levels) - 1

    def sub(m):
        level = max(1, len(m.group(1)) - shift)
        title = re.sub(r"^(?:\d+(?:\.\d+)*|[IVXLC]+)[.)]?\s+", "", m.group(2).strip())
        return "\n" + "#" * level + " " + title + "\n"

    return re.sub(r"^(#{1,6})\s+(.*)$", sub, text, flags=re.M)


def insert_bibliography(text: str, block: str) -> str:
    if not block:
        return text
    m = re.search(r"^#\s+(?:References?|Bibliograph\w*|R[ée]f[ée]rences)\s*$", text, re.M | re.I)
    if not m:
        return text.rstrip() + "\n\n# References" + block
    rest = text[m.end():]
    nxt = re.search(r"^#{1,6}\s", rest, re.M)
    if (rest[: nxt.start()] if nxt else rest).strip():
        return text      # user wrote their own reference list
    return text[: m.end()] + block + rest


def preprocess(raw: str, md_dir: Path, build: Path, target: str, bib_override: Path | None):
    meta, _ = extract_preamble(strip_comments(raw))

    bib = bib_override or (Path(os.path.expanduser(meta["bibliography"])) if meta.get("bibliography") else None)
    if bib is not None and not bib.is_absolute():
        bib = (md_dir / bib).resolve()
    entries: dict[str, dict] = {}
    if bib is not None and not bib.exists():
        warn(f"Bibliography file not found: {bib}")
    elif bib is not None:
        entries = load_bib(bib, build)

    text = strip_comments(raw, keep_figure_markers=True)    # commented-out citations are ignored
    figs = scan_figures(text)
    text = replace_figure_refs(text, figs, target)
    text, keys = replace_citations(text, set(entries), target)
    if not entries and re.search(r"\]\([^()\s/]+\.pdf\)", text):
        warn("Citations found but no usable 'Bibliography File': no reference list.")
    text = build_figures(text, md_dir, build, target)
    text = strip_comments(text)
    text = re.sub(r"\[\[([^\]|]+)\|([^\]]+)\]\]", r"\2", text)     # [[target|alias]]
    text = re.sub(r"\[\[([^\]]+)\]\]", r"\1", text)                # [[target]]
    _, text = extract_preamble(text)
    text = normalise_headings(text)
    text = re.sub(r"(?<![\w\\`])@", r"\\@", text)   # a literal '@' is never a pandoc citation
    text = insert_bibliography(text, bibliography_block(keys, entries, target) if keys else "")
    return meta, text


def pandoc_meta(meta: dict, with_title: bool) -> str:
    def q(s: str) -> str:
        return '"' + s.replace("\\", "\\\\").replace('"', '\\"') + '"'

    lines = ["lang: en-GB"]
    if with_title and meta.get("title"):
        lines.append("title: " + q(meta["title"]))
    return "\n".join(lines) + "\n"


# ==========================================================================
# 2. PDF target (pandoc -> LaTeX -> xelatex)
# ==========================================================================

def tex_escape(s: str) -> str:
    rep = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#",
           "_": r"\_", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\^{}"}
    return "".join(rep.get(c, c) for c in s)


PREAMBLE = r"""
% ---------- ANR AAPG pre-proposal layout (generated by anr_export.py) ----------
\usepackage{fontspec}
\IfFontExistsTF{Calibri}{\setmainfont{Calibri}}{%
  \IfFontExistsTF{Carlito}{\setmainfont{Carlito}}{\typeout{ANR-NOFONT}}}
\definecolor{anrcell}{HTML}{B6DDE8}
\definecolor{anrblue}{HTML}{31849B}
\definecolor{anrpage}{HTML}{4472C4}
\usepackage{array,colortbl}
\usepackage{fancyhdr}
\usepackage{etoolbox}
\usepackage{titlesec}
\usepackage{caption}

% header: the blue table of the official template, on every page
\newcolumntype{A}[1]{>{\columncolor{anrcell}\raggedright\arraybackslash}p{\dimexpr#1\textwidth-2\tabcolsep-1.4pt\relax}}
\newcommand{\anrheader}{%
  {\small\setlength{\arrayrulewidth}{1pt}\arrayrulecolor{white}\renewcommand{\arraystretch}{1.1}%
  \noindent\begin{tabular}{|A{0.2024}|A{0.5794}|A{0.2182}|}\hline
    \textbf{AAPG 2027} & \textbf{<<ACRONYM>>} & <<INSTRUMENT>> \\\hline
    Coordinated by: & <<COORDINATOR>> & <<DURATION>> \\\hline
    \multicolumn{3}{|>{\columncolor{anrcell}\raggedright\arraybackslash}p{\dimexpr\textwidth-2\tabcolsep-2pt\relax}|}{<<THEME>>} \\\hline
  \end{tabular}}}
\pagestyle{fancy}
\fancyhf{}
\renewcommand{\headrulewidth}{0pt}
\fancyhead[C]{\anrheader}
\fancyfoot[R]{\color{anrpage}\footnotesize\thepage}
\fancypagestyle{plain}{\fancyhf{}\fancyhead[C]{\anrheader}\fancyfoot[R]{\color{anrpage}\footnotesize\thepage}}

% headings: I., II. / I.1, I.2 / unnumbered
\renewcommand{\thesection}{\Roman{section}}
\renewcommand{\thesubsection}{\thesection.\arabic{subsection}}
\titleformat{\section}{\color{anrblue}\bfseries\fontsize{14}{17}\selectfont}{\thesection.}{0.6em}{}
\titleformat{\subsection}{\color{anrblue}\bfseries\large}{\thesubsection}{0.6em}{}
\titleformat{\subsubsection}{\color{anrblue}\bfseries\itshape\normalsize}{}{0em}{}
\titleformat{\paragraph}[runin]{\bfseries\itshape\normalsize}{}{0em}{}[.]
\titlespacing*{\section}{0pt}{10pt plus 2pt minus 2pt}{4pt plus 1pt}
\titlespacing*{\subsection}{0pt}{8pt plus 2pt minus 2pt}{3pt plus 1pt}
\titlespacing*{\subsubsection}{0pt}{6pt plus 2pt minus 1pt}{2pt}

% figure captions: "Fig. 2 | Title. Text"
\DeclareCaptionLabelSeparator{anrpipe}{\textbf{\ |\ }}
\captionsetup[figure]{name=Fig., labelfont=bf, labelsep=anrpipe, font=footnotesize,
  justification=justified, singlelinecheck=false, skip=4pt}
\setlength{\textfloatsep}{10pt plus 2pt minus 2pt}
\setlength{\floatsep}{8pt plus 2pt minus 2pt}

% figures: framed floating box; the image keeps its natural size and the
% legend starts beside it, then continues below it (done with \hangindent,
% since wrapfig fails inside floats/boxes)
\usepackage[most]{tcolorbox}
\usepackage{ragged2e}
\newsavebox{\anrimgbox}
\newlength{\anrinnerw}
\newlength{\anrroom}
\newlength{\anrfigminheight}
\newlength{\anrboxpad}\setlength{\anrboxpad}{3pt}
\newlength{\anrfiggap}\setlength{\anrfiggap}{6pt}
\newlength{\anrminroom}\setlength{\anrminroom}{4cm}   % less room than this: legend below
% the box is <<FIGWIDTH>> x the text width, sticking out equally into both margins
\newlength{\anrfigextra}\setlength{\anrfigextra}{\dimexpr(<<FIGWIDTH>>\textwidth-\textwidth)/2\relax}
\newif\ifanrbeside
% #1: file, #2: optional width (empty = natural size)
\newcommand{\anrsetfig}[2]{%
  \setlength{\anrinnerw}{\dimexpr\linewidth+2\anrfigextra-2\anrboxpad-1.2pt\relax}%
  \if\relax\detokenize{#2}\relax
    \sbox{\anrimgbox}{\includegraphics{#1}}%
  \else
    \sbox{\anrimgbox}{\includegraphics[width=#2]{#1}}%
  \fi
  \ifdim\wd\anrimgbox>\anrinnerw
    \sbox{\anrimgbox}{\includegraphics[width=\anrinnerw]{#1}}%
    \PackageWarning{anr}{Figure #1 wider than the text: scaled down}%
  \fi
  \setlength{\anrroom}{\dimexpr\anrinnerw-\wd\anrimgbox-\anrfiggap\relax}%
  \ifdim\anrroom<\anrminroom
    \anrbesidefalse\setlength{\anrfigminheight}{0pt}%
  \else
    \anrbesidetrue
    \setlength{\anrfigminheight}{\dimexpr\ht\anrimgbox+\dp\anrimgbox+2\anrboxpad+1.2pt\relax}%
  \fi}
\newtcolorbox{anrfigbox}[1][]{enhanced, float, floatplacement=htbp,
  colback=white, colframe=anrblue, boxrule=0.6pt, arc=2pt, boxsep=0pt,
  grow to left by=\anrfigextra, grow to right by=\anrfigextra,
  left=\anrboxpad, right=\anrboxpad, top=\anrboxpad, bottom=\anrboxpad,
  before upper={\setlength{\parskip}{0pt}\setlength{\parindent}{0pt}\footnotesize\justifying}, #1}
% number of legend lines running beside the image
\newcommand{\anrfiglines}{\numexpr(\dimexpr\ht\anrimgbox+\dp\anrimgbox+\anrfiggap/2\relax+\baselineskip/2)/\baselineskip\relax}
\newcommand{\anrfigbelow}{{\centering\usebox{\anrimgbox}\par}\vspace{4pt}\noindent}
\newcommand{\anrwrapleft}{%
  \ifanrbeside
    \noindent\hangindent=\dimexpr\wd\anrimgbox+\anrfiggap\relax\hangafter=-\anrfiglines
    \llap{\makebox[0pt][l]{\hspace*{\dimexpr-\wd\anrimgbox-\anrfiggap\relax}%
      \raisebox{\dimexpr\ht\strutbox-\ht\anrimgbox\relax}[0pt][0pt]{\usebox{\anrimgbox}}}}%
  \else\anrfigbelow\fi}
\newcommand{\anrwrapright}{%
  \ifanrbeside
    \noindent\hangindent=-\dimexpr\wd\anrimgbox+\anrfiggap\relax\hangafter=-\anrfiglines
    \rlap{\hspace*{\dimexpr\linewidth-\wd\anrimgbox\relax}%
      \raisebox{\dimexpr\ht\strutbox-\ht\anrimgbox\relax}[0pt][0pt]{\usebox{\anrimgbox}}}%
  \else\anrfigbelow\fi}
\newcommand{\anrfignum}[1]{\refstepcounter{figure}\label{#1}\textbf{Fig.~\thefigure\ |\ }}
\newcommand{\anrfigend}{\par}

% reference list: one running paragraph (link colour set in the text itself)
\newenvironment{anrrefs}{\par\footnotesize\justifying\hypersetup{urlcolor=black}}{\par}

\setlength{\parskip}{4pt plus 1pt minus 1pt}
\setlength{\parindent}{0pt}
\AtBeginDocument{\urlstyle{same}}
"""

TITLE_BLOCK = r"""
\thispagestyle{fancy}
\begin{center}
{\fontsize{14}{17}\selectfont\bfseries <<TITLE>>\par}
\end{center}
\suppressfloats[t]
\vspace{2pt}
"""


def find_tool(names: list[str]) -> str | None:
    extra = ["/Library/TeX/texbin", "/opt/homebrew/bin", "/usr/local/bin"]
    for n in names:
        p = shutil.which(n) or shutil.which(n, path=os.pathsep.join(extra))
        if p:
            return p
    return None


def find_pandoc() -> str:
    p = find_tool(["pandoc"])
    if p:
        return p
    try:
        import pypandoc  # type: ignore
        return pypandoc.get_pandoc_path()
    except Exception:
        sys.exit("pandoc not found: `brew install pandoc` (or `pip install pypandoc_binary`).")


def run_pandoc(args: list[str], cwd: Path) -> None:
    r = subprocess.run([find_pandoc()] + args, capture_output=True, text=True, cwd=cwd)
    if r.returncode != 0:
        sys.exit("pandoc failed:\n" + r.stderr)
    for line in r.stderr.splitlines():
        if line.strip() and "rsvg-convert" not in line:
            warn("pandoc: " + line.strip())


def build_pdf(meta, text, md_dir: Path, build: Path, out: Path) -> Path:
    def field(k, default):
        v = meta.get(k)
        if not v:
            warn(f"Preamble has no '{k}': header shows a placeholder.")
        return tex_escape(v) if v else default

    pre = (PREAMBLE.replace("<<ACRONYM>>", field("acronym", "ACRONYM"))
           .replace("<<INSTRUMENT>>", field("instrument", "Instrument"))
           .replace("<<COORDINATOR>>", field("coordinator", "First name SURNAME"))
           .replace("<<DURATION>>", field("duration", "Duration"))
           .replace("<<THEME>>", field("theme", "Scientific theme")))
    try:
        figwidth = float(meta.get("figwidth", "1.2"))
    except ValueError:
        warn(f"'Figure Width: {meta['figwidth']}' is not a number: using 1.2.")
        figwidth = 1.2
    pre = pre.replace("<<FIGWIDTH>>", f"{figwidth:.3f}")
    margin = meta.get("margin", "2.5cm")
    (build / "anr-preamble.tex").write_text(pre, encoding="utf-8")
    (build / "anr-title.tex").write_text(
        TITLE_BLOCK.replace("<<TITLE>>", field("title", "Title of the project")), encoding="utf-8")
    (build / "meta.yaml").write_text(pandoc_meta(meta, False), encoding="utf-8")
    (build / "body.md").write_text(text, encoding="utf-8")

    tex = build / (out.stem + ".tex")
    args = [str(build / "body.md"), "-f", PANDOC_FORMAT, "-t", "latex", "-s",
            "--metadata-file", str(build / "meta.yaml"),
            "-H", str(build / "anr-preamble.tex"), "-B", str(build / "anr-title.tex"),
            "-V", "documentclass=article", "-V", "fontsize=11pt", "-V", "papersize=a4",
            "-V", f"geometry=a4paper,left={margin},right={margin},bottom={margin},"
                  f"top=\\dimexpr{margin}+0.75cm\\relax,headheight=1.75cm,headsep=0.4cm,footskip=0.9cm",
            "-V", "numbersections", "-V", "secnumdepth=2",
            "-V", "colorlinks", "-V", "linkcolor=black", "-V", "citecolor=black",
            "-V", "urlcolor=anrblue",
            "-V", f"title-meta={meta.get('title', '')}", "-V", f"author-meta={meta.get('coordinator', '')}",
            "-o", str(tex)]
    run_pandoc(args, md_dir)

    # xelatex is called directly (not through latexmk): it writes the PDF itself,
    # whereas latexmk -xelatex stops at an intermediate .xdv file if its
    # xdvipdfmx step fails or is configured differently.
    xelatex = find_tool(["xelatex"])
    if not xelatex:
        sys.exit(f"xelatex not found (install MacTeX or BasicTeX). The LaTeX source is in {tex}")
    pdf = build / (out.stem + ".pdf")
    logf = build / (out.stem + ".log")
    pdf.unlink(missing_ok=True)
    for _ in range(3):   # re-run until cross-references are stable
        r = subprocess.run([xelatex, "-interaction=nonstopmode", "-halt-on-error", tex.name],
                           cwd=build, capture_output=True, text=True)
        log = logf.read_text(errors="replace") if logf.exists() else r.stdout
        if r.returncode or not re.search(r"Rerun to get|Label\(s\) may have changed", log):
            break
    if r.returncode or not pdf.exists():
        errors = "\n".join(l for l in log.splitlines() if l.startswith("!")) or log[-2000:]
        sys.exit(f"LaTeX failed (full log: {logf}):\n{errors}")
    if "ANR-NOFONT" in log:
        warn("Neither Calibri nor Carlito is installed: the PDF uses Latin Modern (serif). "
             "See the README, 'Fonts'.")
    if re.search(r"Reference .* undefined", log):
        warn("LaTeX reports undefined references (see the .log).")
    if "Overfull \\hbox" in log:
        n = log.count("Overfull \\hbox")
        warn(f"{n} overfull line(s) in LaTeX (text sticking into the margin); see the .log.")
    shutil.copy(pdf, out)
    return out


# ==========================================================================
# 3. DOCX target (pandoc -> Word, using the official template)
# ==========================================================================

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}
CAL = '<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:cs="Calibri" w:eastAsia="Calibri"/>'
ACC = "31849B"
DOCX_STYLES = {  # styleId: (name, basedOn, pPr, rPr, next)
    "Normal": ("Normal", None, '<w:widowControl/><w:spacing w:before="0" w:after="0"/><w:jc w:val="left"/>',
               CAL + '<w:sz w:val="22"/><w:szCs w:val="22"/><w:lang w:val="en-GB"/>', None),
    "BodyText": ("Body Text", "Normal", '<w:spacing w:before="0" w:after="100"/><w:jc w:val="both"/>', "", None),
    "FirstParagraph": ("First Paragraph", "BodyText", "", "", "BodyText"),
    "Compact": ("Compact", "BodyText", '<w:spacing w:before="0" w:after="20"/>', "", None),
    "Title": ("Title", "Normal", '<w:spacing w:after="160"/><w:jc w:val="center"/>',
              '<w:b/><w:bCs/><w:sz w:val="28"/><w:szCs w:val="28"/>', "BodyText"),
    "Heading1": ("heading 1", "Normal",
                 '<w:keepNext/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>'
                 '<w:spacing w:before="200" w:after="100"/><w:ind w:left="714" w:hanging="357"/><w:outlineLvl w:val="0"/>',
                 f'<w:b/><w:bCs/><w:color w:val="{ACC}"/><w:sz w:val="28"/><w:szCs w:val="28"/>', "BodyText"),
    "Heading2": ("heading 2", "Normal",
                 '<w:keepNext/><w:numPr><w:ilvl w:val="1"/><w:numId w:val="1"/></w:numPr>'
                 '<w:spacing w:before="160" w:after="80"/><w:ind w:left="567" w:hanging="567"/><w:outlineLvl w:val="1"/>',
                 f'<w:b/><w:bCs/><w:color w:val="{ACC}"/><w:sz w:val="24"/><w:szCs w:val="24"/>', "BodyText"),
    "Heading3": ("heading 3", "Normal", '<w:keepNext/><w:spacing w:before="120" w:after="60"/><w:outlineLvl w:val="2"/>',
                 f'<w:b/><w:bCs/><w:i/><w:iCs/><w:color w:val="{ACC}"/>', "BodyText"),
    "CaptionedFigure": ("Captioned Figure", "Normal",
                        '<w:keepNext/><w:spacing w:before="120" w:after="40"/><w:jc w:val="center"/>', "", None),
    "Figure": ("Figure", "Normal", '<w:keepNext/><w:jc w:val="center"/>', "", None),
    "ImageCaption": ("Image Caption", "Normal", '<w:spacing w:after="160"/><w:jc w:val="both"/>',
                     '<w:sz w:val="18"/><w:szCs w:val="18"/>', None),
    "Bibliography": ("Bibliography", "Normal", '<w:spacing w:after="40"/><w:jc w:val="both"/>',
                     '<w:sz w:val="18"/><w:szCs w:val="18"/>', None),
    "FootnoteText": ("footnote text", "Normal", "", '<w:sz w:val="18"/><w:szCs w:val="18"/>', None),
}
DOCX_PLACEHOLDERS = {
    "ACRONYM": "acronym", "Instrument": "instrument", "First name SURNAME": "coordinator",
    "Duration": "duration",
    "Number and title of the chosen scientific theme (to be found in 2027 AAPG call)": "theme",
}


def build_reference_docx(template: Path, out: Path) -> None:
    import copy
    import zipfile
    from lxml import etree

    with zipfile.ZipFile(template) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    root = etree.fromstring(parts["word/styles.xml"])
    for sid, (name, based, ppr, rpr, nxt) in DOCX_STYLES.items():
        xml = (f'<w:style xmlns:w="{W}" w:type="paragraph" w:styleId="{sid}"><w:name w:val="{name}"/>'
               + (f'<w:basedOn w:val="{based}"/>' if based else "")
               + (f'<w:next w:val="{nxt}"/>' if nxt else "")
               + f'<w:qFormat/><w:pPr>{ppr}</w:pPr><w:rPr>{rpr}</w:rPr></w:style>')
        new = etree.fromstring(xml)
        if based is None:
            new.set(f"{{{W}}}default", "1")
        old = root.find(f'w:style[@w:styleId="{sid}"]', NS)
        if old is not None:
            if old.find("w:link", NS) is not None:
                new.insert(1, copy.deepcopy(old.find("w:link", NS)))
            root.replace(old, new)
        else:
            root.append(new)
    parts["word/styles.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    num = etree.fromstring(parts["word/numbering.xml"])
    ref = num.find('w:num[@w:numId="1"]/w:abstractNumId', NS)
    if ref is not None:
        absn = num.find(f'w:abstractNum[@w:abstractNumId="{ref.get(f"{{{W}}}val")}"]', NS)
        lvl = absn.find('w:lvl[@w:ilvl="1"]', NS) if absn is not None else None
        if lvl is not None:
            absn.replace(lvl, etree.fromstring(
                f'<w:lvl xmlns:w="{W}" w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
                f'<w:lvlText w:val="%1.%2"/><w:lvlJc w:val="left"/>'
                f'<w:pPr><w:ind w:left="567" w:hanging="567"/></w:pPr></w:lvl>'))
    parts["word/numbering.xml"] = etree.tostring(num, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in parts.items():
            z.writestr(n, data)


def fill_docx_header(path: Path, meta: dict) -> None:
    import docx

    doc = docx.Document(str(path))
    for section in doc.sections:
        for hdr in (section.header, section.first_page_header, section.even_page_header):
            if hdr.is_linked_to_previous:
                continue
            for table in hdr.tables:
                seen = set()
                for row in table.rows:
                    for cell in row.cells:
                        key = DOCX_PLACEHOLDERS.get(cell.text.strip())
                        if cell._tc in seen or not key or not meta.get(key):
                            continue
                        seen.add(cell._tc)
                        runs = [r for p in cell.paragraphs for r in p.runs]
                        runs[0].text = meta[key].replace(" -- ", " \u2013 ")
                        for r in runs[1:]:
                            r._r.getparent().remove(r._r)
    doc.core_properties.title = meta.get("title", "")
    doc.core_properties.author = meta.get("coordinator", "")
    doc.save(str(path))


def build_docx(meta, text, md_dir, build, out, template) -> Path:
    ref = build / "reference.docx"
    build_reference_docx(template, ref)
    (build / "meta.yaml").write_text(pandoc_meta(meta, True), encoding="utf-8")
    (build / "body.md").write_text(text, encoding="utf-8")
    args = [str(build / "body.md"), "-f", PANDOC_FORMAT, "-t", "docx", "--reference-doc", str(ref),
            "--metadata-file", str(build / "meta.yaml"), "--dpi=300", "-o", str(out)]
    run_pandoc(args, md_dir)
    fill_docx_header(out, meta)
    return out


# ==========================================================================
# main
# ==========================================================================

def pdf_pages(pdf: Path) -> int:
    info = find_tool(["pdfinfo"])
    if info:
        m = re.search(r"Pages:\s+(\d+)", subprocess.run([info, str(pdf)], capture_output=True, text=True).stdout)
        if m:
            return int(m.group(1))
    return len(re.findall(rb"/Type\s*/Page(?!s)", pdf.read_bytes()))


def export(args) -> None:
    warnings.clear()
    md = args.markdown.resolve()
    target = "docx" if args.docx else "latex"
    out = (args.output or md.with_suffix(".docx" if args.docx else ".pdf")).resolve()
    build = out.parent / f".{out.stem}_build"
    build.mkdir(exist_ok=True)

    meta, text = preprocess(md.read_text(encoding="utf-8"), md.parent, build, target, args.bib)
    if target == "docx":
        build_docx(meta, text, md.parent, build, out, args.template)
        print(f"Written {out}")
    else:
        build_pdf(meta, text, md.parent, build, out)
        n = pdf_pages(out)
        print(f"Written {out}  ({n} page{'s' if n > 1 else ''}"
              f"{', OVER the 4-page limit' if n > 4 else ''})")
    for w_ in warnings:
        print("  ! " + w_)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("markdown", type=Path)
    ap.add_argument("-o", "--output", type=Path, help="output file (default: next to the .md)")
    ap.add_argument("--docx", action="store_true", help="produce the filled Word template instead of a PDF")
    ap.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE, help="ANR .docx template (--docx)")
    ap.add_argument("--bib", type=Path, help="override the preamble's 'Bibliography File'")
    args = ap.parse_args()
    for a in ("bib", "output", "template"):   # paths given on the command line: relative to cwd
        if getattr(args, a) is not None:
            setattr(args, a, getattr(args, a).expanduser().resolve())
    export(args)


if __name__ == "__main__":
    main()